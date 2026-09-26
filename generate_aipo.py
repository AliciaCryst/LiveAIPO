#!/usr/bin/env python3
"""Build AIPO holdings text and an interactive, price-colored treemap."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import json
import logging
import math
from pathlib import Path
import re
import time
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen
from zipfile import is_zipfile

from openpyxl import load_workbook


PAGE_URL = "https://www.defianceetfs.com/aipo-full-holdings/"
DOWNLOAD_API = "https://www.defianceetfs.com/wp-admin/admin-ajax.php"
HERE = Path(__file__).resolve().parent
SYMBOL_PATTERN = re.compile(r"[A-Z][A-Z0-9]{0,9}(?:-[A-Z0-9]{1,3})?\Z")
LOG = logging.getLogger("aipo")


@dataclass(frozen=True)
class Holding:
    ticker: str
    name: str
    weight: float  # percentage points, e.g. 8.1 means 8.1%


def download_holdings() -> bytes:
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            headers = {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
                ),
                "Referer": PAGE_URL,
            }
            lookup = Request(
                DOWNLOAD_API,
                data=urlencode({"slug": "aipo", "action": "download_full_holdings"}).encode(),
                headers={**headers, "Accept": "application/json"},
                method="POST",
            )
            with urlopen(lookup, timeout=45) as response:
                payload = json.loads(response.read())
            file_url = str(payload.get("data", {}).get("url", ""))
            parsed_url = urlparse(file_url)
            if (not payload.get("success") or parsed_url.hostname != "www.defianceetfs.com"
                    or not parsed_url.path.lower().endswith(".xlsx")):
                raise ValueError("Defiance did not return a valid AIPO workbook URL")
            request = Request(file_url, headers={**headers, "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
            with urlopen(request, timeout=45) as response:
                content = response.read()
            # Validate here so an HTML/WAF error cannot replace good output later.
            if not is_zipfile(BytesIO(content)):
                raise ValueError("Defiance did not return an XLSX workbook")
            parse_holdings(content)
            return content
        except Exception as exc:
            last_error = exc
            if attempt < 2:
                time.sleep(2 ** attempt)
    raise RuntimeError("Could not download Defiance AIPO holdings") from last_error


def parse_holdings(content: bytes, min_holdings: int = 50) -> tuple[list[Holding], str, list[tuple[str, str, float]]]:
    """Parse Defiance's AIPO workbook, keeping quoteable equity rows."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, ())
        date_row = next(rows, ())
        if str(first[0] if first else "").strip().upper() != "AIPO":
            raise ValueError("Workbook does not identify itself as AIPO")
        raw_date = date_row[0] if date_row else None
        if isinstance(raw_date, datetime):
            holdings_as_of = raw_date.date().isoformat()
        else:
            try:
                holdings_as_of = datetime.strptime(str(raw_date).strip(), "%m/%d/%Y").date().isoformat()
            except ValueError as exc:
                raise ValueError("Could not read the AIPO holdings date") from exc

        header: dict[str, int] | None = None
        by_ticker: dict[str, Holding] = {}
        excluded: list[tuple[str, str, float]] = []
        for row in rows:
            normalized = [str(value or "").strip().casefold() for value in row]
            if header is None:
                required = {"% of net assets", "name", "identifier"}
                if required.issubset(normalized):
                    header = {key: normalized.index(key) for key in required}
                continue
            name = " ".join(str(row[header["name"]] or "").split())
            original_ticker = str(row[header["identifier"]] or "").strip().upper()
            raw_weight = row[header["% of net assets"]]
            if not name or not original_ticker:
                continue
            try:
                if isinstance(raw_weight, str):
                    weight = float(raw_weight.replace("%", "").replace(",", "").strip())
                else:
                    weight = float(raw_weight)
                    if 0 < weight <= 1:
                        weight *= 100
            except (TypeError, ValueError):
                continue
            if not math.isfinite(weight) or weight <= 0:
                continue
            # Defiance's export includes a money-market sweep and a cash row.
            upper_name = name.upper()
            if ("CASH" in upper_name or "MONEY MARKET" in upper_name
                    or "GOVERNMENT OBLIGATIONS FUND" in upper_name):
                continue
            ticker = original_ticker.replace(".", "-").replace("/", "-")
            if not SYMBOL_PATTERN.fullmatch(ticker):
                excluded.append((name, original_ticker, weight))
                continue
            previous = by_ticker.get(ticker)
            by_ticker[ticker] = Holding(ticker, name, weight + (previous.weight if previous else 0))
        if header is None:
            raise ValueError("Could not find % of Net Assets, Name and Identifier columns")
        result = sorted(by_ticker.values(), key=lambda item: (-item.weight, item.ticker))
        total_weight = sum(item.weight for item in result)
        if len(result) < min_holdings or not 95 < total_weight < 101:
            raise ValueError("Unexpected AIPO holdings count or total weight; workbook format may have changed")
        return result, holdings_as_of, excluded
    finally:
        workbook.close()


def _quote_from_frame(frame, symbol: str) -> dict | None:
    """Read the latest bar and the preceding trading session's final bar."""
    import pandas as pd

    if frame is None or frame.empty:
        return None
    if isinstance(frame.columns, pd.MultiIndex):
        if symbol in frame.columns.get_level_values(0):
            frame = frame[symbol]
        elif symbol in frame.columns.get_level_values(1):
            frame = frame.xs(symbol, axis=1, level=1)
        else:
            return None
    if "Close" not in frame:
        return None
    closes = pd.to_numeric(frame["Close"], errors="coerce")
    closes = closes[(closes > 0) & closes.notna()]
    if len(closes) < 2:
        return None
    stamps = pd.DatetimeIndex(closes.index)
    if stamps.tz is not None:
        stamps = stamps.tz_convert("America/New_York")
    sessions = closes.groupby(stamps.date).last()
    if len(sessions) < 2:
        return None
    last = float(sessions.iloc[-1])
    previous = float(sessions.iloc[-2])
    last_stamp = stamps[-1]
    if stamps.tz is not None and (last_stamp.hour != 0 or last_stamp.minute != 0):
        label = last_stamp.tz_convert("America/Los_Angeles").strftime("%Y-%m-%d %H:%M %Z")
    else:
        label = str(last_stamp.date())
    return {
        "price": round(last, 4),
        "change": round(last - previous, 4),
        "changePct": round((last / previous - 1) * 100, 4),
        "priceDate": label,
    }


def get_quotes(symbols: list[str], batch_size: int = 10, delay: float = 2.0,
               retry_delay: float = 30.0, max_retries: int = 2) -> dict[str, dict]:
    """Pace yfinance batches; back off after an empty batch; use daily fallback."""
    import yfinance as yf

    quotes: dict[str, dict] = {}

    def fetch(batch: list[str], interval: str) -> int:
        try:
            frame = yf.download(
                tickers=batch, period="5d", interval=interval, group_by="ticker",
                auto_adjust=False, prepost=False, threads=2, progress=False, timeout=25,
            )
        except Exception as exc:
            LOG.warning("Yahoo %s batch failed (%s symbols): %s", interval, len(batch), exc)
            return 0
        received = 0
        for symbol in batch:
            quote = _quote_from_frame(frame, symbol)
            if quote:
                quotes[symbol] = quote
                received += 1
        return received

    for start in range(0, len(symbols), batch_size):
        batch = symbols[start:start + batch_size]
        if start:
            time.sleep(delay)
        received = fetch(batch, "1m")
        if not received:
            for attempt in range(max_retries):
                wait = retry_delay * 2 ** attempt
                LOG.warning("No prices for %s batch; retrying after %.0f seconds", "1m", wait)
                time.sleep(wait)
                received = fetch(batch, "1m")
                if received:
                    break
            if not received:
                raise RuntimeError("Yahoo returned no prices after paced retries; existing output was kept")
    missing = [symbol for symbol in symbols if symbol not in quotes]
    # Limit fallback requests when Yahoo has blocked the whole run.
    if len(missing) <= len(symbols) // 2:
        for start in range(0, len(missing), batch_size):
            time.sleep(delay)
            fetch(missing[start:start + batch_size], "1d")
    missing = [symbol for symbol in symbols if symbol not in quotes]
    LOG.info("Prices received: %d/%d", len(quotes), len(symbols))
    if len(missing) > max(10, len(symbols) // 10):
        raise RuntimeError(f"Yahoo returned too few quotes ({len(quotes)}/{len(symbols)}); existing output was kept")
    if missing:
        LOG.warning("Missing Yahoo quotes (shown as N/A): %s", ", ".join(missing))
    return quotes


def create_outputs(holdings: list[Holding], as_of: str, quotes: dict[str, dict], output_dir: Path) -> None:
    text_lines = [f"# AIPO holdings as of {as_of}; weight_pct is percentage points",
                  "symbol\tcompany\tweight_pct"]
    records = []
    for rank, item in enumerate(holdings, 1):
        text_lines.append(f"{item.ticker}\t{item.name}\t{item.weight:.6f}")
        records.append({"ticker": item.ticker, "name": item.name,
                        "weight": round(item.weight, 6), "rank": rank,
                        **quotes.get(item.ticker, {"price": None, "change": None,
                                                    "changePct": None, "priceDate": None})})
    meta = {"holdingsAsOf": as_of, "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            "totalWeight": round(sum(item.weight for item in holdings), 6),
            "missingQuotes": sum(item.ticker not in quotes for item in holdings)}
    # Prevent a company name from closing the inline script if source text ever changes.
    def safe_json(value: object) -> str:
        return json.dumps(value, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026")

    template = (HERE / "treemap_template.html").read_text(encoding="utf-8")
    html = template.replace("__HOLDINGS_JSON__", safe_json(records)).replace("__METADATA_JSON__", safe_json(meta))
    output_dir.mkdir(parents=True, exist_ok=True)
    # Construct both results before touching any existing output.
    text_path = output_dir / "aipo.txt"
    html_path = output_dir / "aipo.html"
    text_path.with_suffix(".txt.tmp").write_text("\n".join(text_lines) + "\n", encoding="utf-8")
    html_path.with_suffix(".html.tmp").write_text(html, encoding="utf-8")
    text_path.with_suffix(".txt.tmp").replace(text_path)
    html_path.with_suffix(".html.tmp").replace(html_path)
    LOG.info("Wrote %d holdings (%.6f%%) to %s", len(holdings), meta["totalWeight"], output_dir)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Local Defiance AIPO XLSX for testing; default downloads it")
    parser.add_argument("--output-dir", type=Path, default=HERE)
    parser.add_argument("--quote-batch-size", type=int, default=10)
    parser.add_argument("--quote-delay", type=float, default=2.0, help="Seconds between Yahoo batches")
    parser.add_argument("--quote-retry-delay", type=float, default=30.0, help="Initial backoff after an empty batch")
    args = parser.parse_args()
    if args.quote_batch_size < 1 or args.quote_delay < 0 or args.quote_retry_delay < 0:
        parser.error("Batch size must be positive and delays must be nonnegative")
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    content = args.source.read_bytes() if args.source else download_holdings()
    holdings, as_of, excluded = parse_holdings(content)
    for name, symbol, weight in excluded:
        LOG.warning("Excluded unquotable row: %s (%s), %.6f%%", name, symbol, weight)
    quotes = get_quotes([holding.ticker for holding in holdings],
                        batch_size=args.quote_batch_size, delay=args.quote_delay,
                        retry_delay=args.quote_retry_delay)
    create_outputs(holdings, as_of, quotes, args.output_dir)


if __name__ == "__main__":
    main()
